"""上传文件解析器（B-304）。

支持 CSV / JSON / XLSX / TXT，统一输出 ``ParsedUpload``：
- ``records``：结构化行（list[dict]），供后续清洗 Agent 使用；
- ``raw_text``：文本原文（xlsx 为二进制格式，raw_text 为空）；
- 解析失败抛 ``ParseError``，由 API 层映射为 415 ``FILE_FORMAT_INVALID``。
"""
from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

ALLOWED_EXTENSIONS = {"csv", "json", "xlsx", "txt"}
_TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030")


class ParseError(ValueError):
    """文件解析失败（格式/编码/内容非法）。"""


@dataclass(frozen=True)
class ParsedUpload:
    format: str
    encoding: str | None
    records: list[dict[str, Any]] = field(default_factory=list)
    raw_text: str = ""

    @property
    def record_count(self) -> int:
        return len(self.records)


def _decode(content: bytes) -> tuple[str, str]:
    for encoding in _TEXT_ENCODINGS:
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ParseError("无法识别文件编码（支持 UTF-8 / GB18030）")


def _as_dict(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    return {"value": value}


def parse_upload(content: bytes, filename: str) -> ParsedUpload:
    """按扩展名分派解析器。"""
    suffix = Path(filename).suffix.lower()
    if not content:
        raise ParseError("文件内容为空")

    if suffix == ".csv":
        return _parse_csv(content)
    if suffix == ".json":
        return _parse_json(content)
    if suffix == ".xlsx":
        return _parse_xlsx(content)
    if suffix == ".txt":
        return _parse_txt(content)
    raise ParseError(f"不支持的文件格式：{suffix or '<无扩展名>'}")


def _parse_csv(content: bytes) -> ParsedUpload:
    text, encoding = _decode(content)
    reader = csv.DictReader(StringIO(text))
    records = [
        {str(key): value for key, value in row.items() if key is not None}
        for row in reader
        if any(value for value in row.values())
    ]
    if not records:
        raise ParseError("CSV 没有可解析的数据行（需包含表头）")
    return ParsedUpload(format="csv", encoding=encoding, records=records, raw_text=text)


def _parse_json(content: bytes) -> ParsedUpload:
    text, encoding = _decode(content)
    try:
        data = json.loads(text)
    except ValueError as exc:
        raise ParseError(f"JSON 解析失败：{exc}") from exc

    if isinstance(data, list):
        records = [_as_dict(item) for item in data]
    elif isinstance(data, dict):
        nested = data.get("records") or data.get("data")
        if isinstance(nested, list):
            records = [_as_dict(item) for item in nested]
        else:
            records = [data]
    else:
        records = [{"value": data}]

    if not records:
        raise ParseError("JSON 内容为空")
    return ParsedUpload(format="json", encoding=encoding, records=records, raw_text=text)


def _parse_xlsx(content: bytes) -> ParsedUpload:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl 异常类型多样
        raise ParseError(f"XLSX 解析失败：{exc}") from exc

    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise ParseError("XLSX 首个工作表为空")

    header = [str(cell) if cell is not None else f"col_{i}" for i, cell in enumerate(rows[0])]
    records: list[dict[str, Any]] = []
    for row in rows[1:]:
        if row is None or all(cell is None for cell in row):
            continue
        records.append(
            {header[i] if i < len(header) else f"col_{i}": cell for i, cell in enumerate(row)}
        )
    if not records:
        raise ParseError("XLSX 没有可解析的数据行")
    return ParsedUpload(format="xlsx", encoding=None, records=records, raw_text="")


def _parse_txt(content: bytes) -> ParsedUpload:
    text, encoding = _decode(content)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        raise ParseError("TXT 内容为空")
    records = [{"text": line} for line in lines]
    return ParsedUpload(format="txt", encoding=encoding, records=records, raw_text=text)


__all__ = [
    "ALLOWED_EXTENSIONS",
    "ParseError",
    "ParsedUpload",
    "parse_upload",
]
