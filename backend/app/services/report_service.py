"""报告服务（B-508 / B-509 / B-510）。

- 生成：建 ``reports``（generating）→ Celery ``tasks.generate_report`` → ReportGenerationAgent
  → Markdown/Excel 渲染 → 存对象存储 → completed（失败 failed）；
- 查询：列表（类型/状态过滤 + 分页）、详情（累加 view_count）；
- 下载：读取文件字节 + 累加 download_count；
- 删除：硬删除（API 204）；
- 格式口径（D-11 降级）：MVP 支持 markdown / excel；pdf / html 返回 422 REPORT_FORMAT_INVALID。
"""
from __future__ import annotations

import time
import uuid
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from loguru import logger
from sqlalchemy import Float, case, cast, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import assert_owner
from app.api.pagination import PaginationParams, paginate
from app.core.exceptions import (
    DataSourceNotFoundError,
    ReportFormatInvalidError,
    ReportNotFoundError,
)
from app.models import AnalysisResult, CleanedData, DataSource, Report, Task, User
from app.schemas.agents.report import (
    REPORT_TYPE_LABELS,
    REPORT_TYPES,
    SUPPORTED_FORMATS,
    ReportConfig,
    ReportInput,
)
from app.services import knowledge_status, storage_service

TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "templates" / "reports"
MEDIA_TYPES = {
    "markdown": "text/markdown; charset=utf-8",
    "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
EXTENSIONS = {"markdown": "md", "excel": "xlsx"}
_SORTABLE = {"created_at": Report.created_at, "completed_at": Report.completed_at}


def _now_naive() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _build_agent():  # noqa: ANN202 - 允许测试 monkeypatch
    from app.agents.report_agent import ReportGenerationAgent

    return ReportGenerationAgent()


def enqueue_report(report: Report) -> str:
    from app.tasks.report import generate_report_task

    return str(generate_report_task.delay(str(report.id)))


# --------------------------------------------------------------------------- 数据聚合


async def _quality_data(session: AsyncSession, datasource_id: uuid.UUID) -> dict:
    overall = cast(AnalysisResult.result["overall_score"].astext, Float)
    row = (
        await session.execute(
            select(
                func.count().label("total"),
                func.coalesce(func.avg(overall), 0).label("avg_score"),
                func.coalesce(
                    func.sum(
                        case(
                            (
                                AnalysisResult.result["needs_review"].astext == "true",
                                1,
                            ),
                            else_=0,
                        )
                    ),
                    0,
                ).label("needs_review"),
                *[
                    func.coalesce(
                        func.avg(
                            cast(
                                AnalysisResult.result["dimensions"][dimension]["score"].astext,
                                Float,
                            )
                        ),
                        0,
                    ).label(f"avg_{dimension}")
                    for dimension in (
                        "accuracy",
                        "completeness",
                        "hallucination",
                        "tone",
                        "compliance",
                    )
                ],
            ).where(
                AnalysisResult.datasource_id == datasource_id,
                AnalysisResult.analysis_type == "quality",
            )
        )
    ).one()
    return {
        "total": int(row.total or 0),
        "avg_overall_score": round(float(row.avg_score or 0), 2),
        "needs_review_count": int(row.needs_review or 0),
        "dimension_averages": {
            dimension: round(float(getattr(row, f"avg_{dimension}") or 0), 2)
            for dimension in (
                "accuracy",
                "completeness",
                "hallucination",
                "tone",
                "compliance",
            )
        },
    }


async def _knowledge_data(session: AsyncSession, datasource_id: uuid.UUID) -> dict:
    status = await knowledge_status.get_knowledge_status(session, datasource_id)
    latest_task = await session.scalar(
        select(Task)
        .where(
            Task.datasource_id == datasource_id,
            Task.task_type == "extract",
            Task.status == "completed",
        )
        .order_by(Task.completed_at.desc())
        .limit(1)
    )
    gaps = ((latest_task.result or {}).get("gaps") if latest_task else []) or []
    return {
        **status,
        "gaps_count": len(gaps),
        "top_gaps": gaps[:5],
    }


async def _insight_data(session: AsyncSession, datasource_id: uuid.UUID) -> dict | None:
    total = int(
        await session.scalar(
            select(func.count())
            .select_from(AnalysisResult)
            .where(
                AnalysisResult.datasource_id == datasource_id,
                AnalysisResult.analysis_type == "insight",
            )
        )
        or 0
    )
    if total == 0:
        return None
    latest = await session.scalar(
        select(AnalysisResult)
        .where(
            AnalysisResult.datasource_id == datasource_id,
            AnalysisResult.analysis_type == "insight",
        )
        .order_by(AnalysisResult.created_at.desc())
        .limit(1)
    )
    return {"total": total, "latest": (latest.result if latest else None)}


async def _collect_data(session: AsyncSession, report: Report) -> dict[str, Any]:
    data: dict[str, Any] = {}
    if report.report_type in ("quality", "comprehensive"):
        data["quality"] = await _quality_data(session, report.datasource_id)
    if report.report_type in ("knowledge", "comprehensive"):
        data["knowledge"] = await _knowledge_data(session, report.datasource_id)
    if report.report_type in ("insight", "comprehensive"):
        data["insight"] = await _insight_data(session, report.datasource_id)
    data["dialogue_count"] = int(
        await session.scalar(
            select(func.count())
            .select_from(CleanedData)
            .where(CleanedData.datasource_id == report.datasource_id)
        )
        or 0
    )
    return data


# --------------------------------------------------------------------------- 渲染


def _bullets(items: list[str]) -> str:
    return "\n".join(f"- {item}" for item in items) if items else "- 无"


def _sections_markdown(sections: list[dict[str, Any]]) -> str:
    blocks = []
    for section in sections:
        blocks.append(f"## {section.get('title', '章节')}\n\n{section.get('content', '')}")
    return "\n\n".join(blocks)


def _quality_status_text(data: dict[str, Any]) -> str:
    quality = data.get("quality") or {}
    if not quality or not quality.get("total"):
        return "暂无质量评估数据"
    dims = "、".join(
        f"{name} {score}" for name, score in (quality.get("dimension_averages") or {}).items()
    )
    return (
        f"- 评估条数：{quality.get('total', 0)}\n"
        f"- 平均总分：{quality.get('avg_overall_score', 0)}\n"
        f"- 需复核：{quality.get('needs_review_count', 0)}\n"
        f"- 维度均分：{dims}"
    )


def _knowledge_status_text(data: dict[str, Any]) -> str:
    knowledge = data.get("knowledge") or {}
    if not knowledge:
        return "暂无知识库数据"
    return (
        f"- 知识总量：{knowledge.get('total', 0)}（种子 {knowledge.get('seed', 0)} / "
        f"提取 {knowledge.get('extracted', 0)} / 手动 {knowledge.get('manual', 0)}）\n"
        f"- 冷启动阶段：{knowledge.get('stage', 'unknown')}\n"
        f"- 知识盲区：{knowledge.get('gaps_count', 0)} 条"
    )


def render_markdown(report: Report, output: dict[str, Any], data: dict[str, Any]) -> str:
    template_path = TEMPLATE_DIR / f"{report.report_type}.md"
    template = template_path.read_text(encoding="utf-8")
    generated_at = output.get("generated_at") or _now_naive().isoformat()
    status = output.get("metadata") or {}
    insight_missing = bool(status.get("insight_degraded"))
    time_range = (report.config or {}).get("time_range") or {}
    time_range_text = (
        f"{time_range.get('start', '')} ~ {time_range.get('end', '')}" if time_range else "未指定"
    )
    replacements = {
        "__TITLE__": report.title,
        "__REPORT_TYPE_LABEL__": REPORT_TYPE_LABELS.get(report.report_type, report.report_type),
        "__GENERATED_AT__": str(generated_at),
        "__TIME_RANGE__": time_range_text,
        "__SUMMARY__": output.get("summary", ""),
        "__HIGHLIGHTS__": _bullets(output.get("highlights") or []),
        "__SECTIONS__": _sections_markdown(output.get("sections") or []),
        "__RECOMMENDATIONS__": _bullets(output.get("recommendations") or []),
        "__QUALITY_STATUS__": _quality_status_text(data),
        "__KNOWLEDGE_STATUS__": _knowledge_status_text(data),
        "__INSIGHT_NOTICE__": ("暂无洞察数据（报告已降级生成）" if insight_missing else "洞察数据已包含在报告中"),
    }
    for token, value in replacements.items():
        template = template.replace(token, str(value))
    return template


def render_excel(report: Report, output: dict[str, Any]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["属性", "值"])
    summary.append(["标题", report.title])
    summary.append(["报告类型", REPORT_TYPE_LABELS.get(report.report_type, report.report_type)])
    summary.append(["生成时间", str(output.get("generated_at", ""))])
    summary.append(["摘要", output.get("summary", "")])
    summary.append(["核心发现", "\n".join(output.get("highlights") or [])])

    sections = workbook.create_sheet("Sections")
    sections.append(["章节", "内容"])
    for section in output.get("sections") or []:
        sections.append([section.get("title", ""), section.get("content", "")])

    recommendations = workbook.create_sheet("Recommendations")
    recommendations.append(["序号", "建议"])
    for index, item in enumerate(output.get("recommendations") or [], start=1):
        recommendations.append([index, item])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- 业务


def validate_create(*, report_type: str, format: str) -> None:
    if report_type not in REPORT_TYPES:
        raise ReportFormatInvalidError(
            details={"report_type": report_type, "allowed": list(REPORT_TYPES)}
        )
    # D-11：MVP 不生成 pdf/html（weasyprint 依赖未引入）
    if format not in SUPPORTED_FORMATS:
        raise ReportFormatInvalidError(
            details={"format": format, "allowed": list(SUPPORTED_FORMATS)}
        )


async def _datasource_or_404(
    session: AsyncSession, datasource_id: uuid.UUID, user: User
) -> DataSource:
    datasource = await session.get(DataSource, datasource_id)
    if datasource is None or datasource.deleted_at is not None:
        raise DataSourceNotFoundError(details={"datasource_id": str(datasource_id)})
    assert_owner(datasource.user_id, user)
    return datasource


async def create_report(
    session: AsyncSession,
    *,
    user: User,
    datasource_id: uuid.UUID,
    title: str,
    report_type: str,
    format: str,
    config: ReportConfig,
) -> Report:
    await _datasource_or_404(session, datasource_id, user)
    validate_create(report_type=report_type, format=format)

    report = Report(
        datasource_id=datasource_id,
        user_id=user.id,
        title=title,
        report_type=report_type,
        format=format,
        status="generating",
        config=config.model_dump(mode="json"),
    )
    session.add(report)
    await session.flush()
    await session.refresh(report)
    return report


async def _execute_report(
    report_id: uuid.UUID, session: AsyncSession, *, commit: bool
) -> dict[str, Any]:
    report = await session.get(Report, report_id)
    if report is None:
        return {"status": "missing", "report_id": str(report_id)}

    started = time.perf_counter()
    report.status = "generating"
    report.error_message = None
    await session.flush()

    try:
        data = await _collect_data(session, report)
        config = ReportConfig(**(report.config or {}))
        payload = ReportInput(
            datasource_id=report.datasource_id,
            title=report.title,
            report_type=report.report_type,  # type: ignore[arg-type]
            data=data,
            config=config,
        )
        agent = _build_agent()
        result = await agent.generate(
            payload.model_dump(mode="json"),
            session=session,
            user_id=report.user_id,
            datasource_id=report.datasource_id,
        )
        if not result.success or not result.data:
            raise RuntimeError("report generation failed: " + "; ".join(result.errors))

        output = result.data
        if report.format == "markdown":
            content = render_markdown(report, output, data).encode("utf-8")
        else:
            content = render_excel(report, output)
        extension = EXTENSIONS[report.format]
        object_name = f"reports/{report.id}.{extension}"
        stored = storage_service.save_bytes(
            content,
            object_name=object_name,
            content_type=MEDIA_TYPES[report.format],
        )

        now = _now_naive()
        report.status = "completed"
        report.file_path = stored.object_name
        report.file_size = stored.size
        report.data = output
        report.generation_duration = round(time.perf_counter() - started, 3)
        report.completed_at = now
        report.expires_at = now + timedelta(days=30)
        await session.flush()
        if commit:
            await session.commit()
        return {
            "status": "completed",
            "report_id": str(report.id),
            "file_size": stored.size,
        }
    except Exception as exc:  # noqa: BLE001
        logger.opt(exception=exc).error("report generation failed report={}", report.id)
        report.status = "failed"
        report.error_message = str(exc)[:1000]
        report.generation_duration = round(time.perf_counter() - started, 3)
        report.completed_at = _now_naive()
        await session.flush()
        if commit:
            await session.commit()
        return {"status": "failed", "report_id": str(report.id), "error": str(exc)}


async def run_report_task(report_id: str, *, session: AsyncSession | None = None) -> dict[str, Any]:
    if session is not None:
        return await _execute_report(uuid.UUID(report_id), session, commit=False)

    from app.core.database import SessionLocal

    async with SessionLocal() as db:
        result = await _execute_report(uuid.UUID(report_id), db, commit=True)
    from app.core.database import dispose_engine

    await dispose_engine()
    return result


async def list_reports(
    session: AsyncSession,
    *,
    user: User,
    datasource_id: uuid.UUID,
    params: PaginationParams,
    report_type: str | None = None,
    status: str | None = None,
) -> tuple[list[Report], Any]:
    await _datasource_or_404(session, datasource_id, user)
    stmt = select(Report).where(Report.datasource_id == datasource_id)
    if report_type:
        stmt = stmt.where(Report.report_type == report_type)
    if status:
        stmt = stmt.where(Report.status == status)
    return await paginate(session, stmt, params, sortable=_SORTABLE)


async def _report_or_404(session: AsyncSession, report_id: uuid.UUID, user: User) -> Report:
    report = await session.get(Report, report_id)
    if report is None:
        raise ReportNotFoundError(details={"report_id": str(report_id)})
    await _datasource_or_404(session, report.datasource_id, user)
    return report


async def get_report(session: AsyncSession, *, user: User, report_id: uuid.UUID) -> Report:
    report = await _report_or_404(session, report_id, user)
    report.view_count = int(report.view_count or 0) + 1
    await session.flush()
    await session.refresh(report)
    return report


async def download_report(
    session: AsyncSession, *, user: User, report_id: uuid.UUID
) -> tuple[str, str, bytes]:
    report = await _report_or_404(session, report_id, user)
    if report.status != "completed" or not report.file_path:
        raise ReportNotFoundError(details={"report_id": str(report_id), "status": report.status})
    report.download_count = int(report.download_count or 0) + 1
    await session.flush()
    content = storage_service.read_bytes(report.file_path)
    extension = EXTENSIONS[report.format]
    filename = f"{report.id}.{extension}"
    return filename, MEDIA_TYPES[report.format], content


async def delete_report(session: AsyncSession, *, user: User, report_id: uuid.UUID) -> None:
    report = await _report_or_404(session, report_id, user)
    await session.delete(report)
    await session.flush()


__all__ = [
    "MEDIA_TYPES",
    "TEMPLATE_DIR",
    "create_report",
    "delete_report",
    "download_report",
    "enqueue_report",
    "get_report",
    "list_reports",
    "render_excel",
    "render_markdown",
    "run_report_task",
    "validate_create",
]
